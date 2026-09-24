"""Unit tests for export_document PDF / PNG / pack formats."""

from __future__ import annotations

from pathlib import Path

import pytest
from django.test import override_settings

from ai.plugins.export_document import ExportDocument


@pytest.fixture
def plugin():
    return ExportDocument()


@pytest.mark.asyncio
async def test_export_pack_writes_four_formats(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Board Pack Demo",
                "format": "pack",
                "content": (
                    "## Summary\n"
                    "- Headcount up across GOFSCO sites this quarter.\n"
                    "- GOSI employer contributions remain within statutory bands.\n"
                    "- No material variance vs last committed payroll cycle.\n"
                ),
                "table": {
                    "headers": ["Metric", "Value"],
                    "rows": [["Headcount", "529"], ["GOSI", "184220"]],
                },
            },
            ctx=None,
        )

    assert "error" not in result, result
    formats = {f["format"] for f in result["files"]}
    assert formats == {"docx", "xlsx", "pdf", "png"}
    out = Path(tmp_path) / "ai_exports"
    for entry in result["files"]:
        path = out / entry["filename"]
        assert path.exists()
        assert path.stat().st_size > 40
    pdf = next(f for f in result["files"] if f["format"] == "pdf")
    assert (out / pdf["filename"]).read_bytes()[:5] == b"%PDF-"
    png = next(f for f in result["files"] if f["format"] == "png")
    assert (out / png["filename"]).read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


@pytest.mark.asyncio
async def test_export_pdf_only(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "PDF Only",
                "format": "pdf",
                "content": (
                    "## Compliance note\n"
                    "October payroll figures were validated against GOSI and KL "
                    "statutory rates. Employer share matches the published schedule "
                    "for the period under review."
                ),
            },
            ctx=None,
        )
    assert "error" not in result
    assert [f["format"] for f in result["files"]] == ["pdf"]


@pytest.mark.asyncio
async def test_export_png_fallback_without_table(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Chart Fallback",
                "format": "png",
                "content": (
                    "## Findings\n"
                    "- Alpha cohort average base pay rose 3.2% versus prior period.\n"
                    "- Beta sites remain within GOSI contribution tolerance bands.\n"
                ),
            },
            ctx=None,
        )
    assert "error" not in result
    assert result["files"][0]["format"] == "png"


@pytest.mark.asyncio
async def test_export_refuses_hollow_placeholder(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Empty Report",
                "format": "docx",
                "content": "[Placeholder for insights]\n[Chart and table to be inserted]",
            },
            ctx=None,
        )
    assert "error" in result
    assert "refused" in result["error"].lower() or "findings" in result["error"].lower()


@pytest.mark.asyncio
async def test_export_refuses_mid_run_prose(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Partial Pack",
                "format": "docx",
                "content": (
                    "## Status\n"
                    "Analysis is still in progress. Partial results will follow "
                    "once the payroll computation completes and validation finishes."
                ),
            },
            ctx=None,
        )
    assert "error" in result
    assert "in-progress" in result["error"].lower() or "pending" in result["error"].lower()


@pytest.mark.asyncio
async def test_export_refuses_thin_prose(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Thin",
                "format": "docx",
                "content": "Looks fine.",
            },
            ctx=None,
        )
    assert "error" in result
    assert "thin" in result["error"].lower() or "refused" in result["error"].lower()


@pytest.mark.asyncio
async def test_export_refuses_insert_slot_shell(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Salary Distribution Analysis at GOFSCO",
                "format": "docx",
                "content": (
                    "## Overview\n"
                    "This report analyses salary distribution.\n\n"
                    "## Key Findings\n"
                    "- [Insert specific insights, e.g., 'Kuwaiti nationals…']\n"
                ),
                "table": {
                    "headers": ["Category", "Average Salary"],
                    "rows": [["Kuwaiti", "[Avg Kuwaiti Salary]"]],
                },
            },
            ctx=None,
        )
    assert "error" in result
    assert "template" in result["error"].lower() or "unfilled" in result["error"].lower()


@pytest.mark.asyncio
async def test_export_docx_renders_gfm_tables_as_word_tables(plugin, tmp_path):
    """GFM pipe tables in ``content`` become real Word tables — never literal |---|."""
    content = (
        "## القروض المفتوحة\n\n"
        "| نوع القرض | المبلغ الأساسي | سعر الفائدة | المدة (شهر) | الحالة |\n"
        "|---|---|---|---|---|\n"
        "| قرض شخصي | 547,000 | 12% | 12 | قيد الإعداد |\n\n"
        "## رصيد الإجازات\n\n"
        "| النوع | المتبقي | المستحق |\n"
        "|---|---|---|\n"
        "| annual | 19 | 30 |\n"
        "| sick | 15 | 15 |\n\n"
        "## ملخص الوضع\n"
        "- قرض شخصي قيد الإعداد بمبلغ 547,000.\n"
        "- رصيد الإجازة السنوية المتبقي 19 يوماً.\n"
    )
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "تقرير القروض المفتوحة ورصيد الإجازات",
                "format": "docx",
                "content": content,
            },
            ctx=None,
        )
    assert "error" not in result, result
    from docx import Document

    path = Path(tmp_path) / "ai_exports" / result["files"][0]["filename"]
    doc = Document(str(path))
    para_text = "\n".join(p.text for p in doc.paragraphs)
    assert "|---|" not in para_text
    assert "| نوع القرض |" not in para_text
    assert "| annual |" not in para_text
    assert len(doc.tables) >= 2
    first = doc.tables[0]
    assert "نوع القرض" in first.rows[0].cells[0].text
    assert "قرض شخصي" in first.rows[1].cells[0].text
    assert "547,000" in first.rows[1].cells[1].text
    second = doc.tables[1]
    assert "annual" in second.rows[1].cells[0].text
    assert "19" in second.rows[1].cells[1].text


@pytest.mark.asyncio
async def test_export_docx_identity_header(plugin, tmp_path):
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "Identity Check",
                "format": "docx",
                "content": (
                    "## Findings\n"
                    "Kuwaiti average base pay is 1,180 KWD across the sampled cohort. "
                    "Indian and Egyptian bands sit between 820 and 980 KWD for the same period."
                ),
                "table": {
                    "headers": ["Nationality", "Avg"],
                    "rows": [["Kuwaiti", "1180"], ["Indian", "900"]],
                },
            },
            ctx=None,
        )
    assert "error" not in result, result
    from docx import Document

    path = Path(tmp_path) / "ai_exports" / result["files"][0]["filename"]
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert any("Pulse — AI Coworker" in t for t in texts)
    assert not any("Carbon" in t for t in texts)


@pytest.mark.asyncio
async def test_export_docx_embeds_image(plugin, tmp_path):
    # Minimal 1x1 PNG
    import base64

    png = base64.b64encode(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    ).decode("ascii")
    with override_settings(MEDIA_ROOT=str(tmp_path)):
        result = await plugin.execute(
            {
                "title": "With Chart",
                "format": "docx",
                "content": "## Findings\n- Real measured result.",
                "images": [{"caption": "Salary chart", "image_b64": png}],
            },
            ctx=None,
        )
    assert "error" not in result, result
    path = Path(tmp_path) / "ai_exports" / result["files"][0]["filename"]
    assert path.exists()
    assert path.read_bytes()[:2] == b"PK"

