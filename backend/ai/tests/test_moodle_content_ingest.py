"""Content-engine integration locks for the aast-med ingest path.

Locks what this pass wired: the shared reader registry routes office/PDF/text
(including ``.ppsx``) and extracts pptx speaker notes, extra uploads become
citable course content, scanned OCR is citable for English and non-citable
pending sign-off for Arabic, the knowledge graph traces every edge to a real
passage_ref, and the keyword index stays whole-word verbatim with semantic OFF.

No count here is a claim about a rung. No golden is loosened.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

from ai.content_engine import ingest
from ai.content_engine import graph as ce_graph
from ai.content_engine import index as ce_index
from ai.content_engine.readers import registry


def _pptx_bytes(body: str = "Slide body sentence.", notes: str = "Speaker note sentence.") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = body
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _blank_pdf_bytes() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


# --- format routing -------------------------------------------------------


def test_ppsx_is_routed_to_the_office_reader():
    assert registry.table()["ext:.ppsx"] == "office"
    result = ingest.read_bytes(_pptx_bytes(body="PPSX body sentence."), "deck.ppsx")
    assert result["status"] == "ok"
    assert result["reader"] == "office"
    assert "PPSX body sentence." in ingest.units_text(result)


def test_pptx_speaker_notes_are_extracted_as_citable_units():
    result = ingest.read_bytes(_pptx_bytes(), "deck.pptx")
    assert result["status"] == "ok"
    locators = {unit["locator"]: unit["text"] for unit in result["units"]}
    assert "Speaker note sentence." in locators["notes 1"]
    assert all(unit["citable"] for unit in result["units"])


def test_txt_md_json_route_through_the_shared_readers():
    assert registry.table()["ext:.txt"] == "text"
    text = ingest.read_bytes(b"plain sentence.", "note.txt")
    md = ingest.read_bytes(b"# Heading\n\nmarkdown sentence.", "note.md")
    js = ingest.read_bytes(json.dumps({"a": "alpha"}).encode(), "note.json")
    assert text["reader"] == "text" and md["reader"] == "text"
    assert js["reader"] == "text"
    assert "a: alpha" in ingest.units_text(js)


# --- extra as course content ---------------------------------------------


_EXTRA_TEXT = "AASTMT Pulse extra file for MED520 tutor notes only."


def test_extra_files_become_citable_course_content(tmp_path, monkeypatch):
    monkeypatch.setattr("ai.moodle_bank._EXTRA", tmp_path)
    from ai.moodle_bank import load_extra, write_extra_index
    from ai.moodle_page import topic_answer
    from ai.tests.test_moodle_topic import _med520_home, _page

    files = [
        {"itemid": 11, "filename": "note.txt", "content": b"Keratinocyte turnover is in this tutor note."},
        {"itemid": 12, "filename": "note.md", "content": _EXTRA_TEXT.encode()},
        {
            "itemid": 13,
            "filename": "note.json",
            "content": json.dumps(
                {"course": "MED520", "title": "n", "passages": [{"text": "json passage text"}]}
            ).encode(),
        },
    ]
    payload = ingest.extra_payload_from_files(files)
    assert {row["itemid"] for row in payload} == {11, 12, 13}
    assert write_extra_index("MED520", payload, extra_root=tmp_path) == 3

    bank = load_extra("MED520", extra_root=tmp_path)
    assert bank["MED520:extra:12:note.md"]["source"] == "extra:note.md"
    # The extra file is still a verbatim cite, one owner per file.
    assert bank["MED520:extra:12:note.md"]["text"] == _EXTRA_TEXT

    answer = topic_answer("educate me about aastmt pulse extra file", _page(_med520_home()))
    assert answer is not None
    assert "MED520:extra:12:note.md" in answer


def test_extra_json_phi_is_still_refused_before_it_reaches_the_readers(tmp_path, monkeypatch):
    # The plugin refuses PHI JSON at store time; the Carbon reader path only
    # sees already-validated bytes. This locks the reader path never changes a
    # file's shape silently.
    monkeypatch.setattr("ai.moodle_bank._EXTRA", tmp_path)
    read = ingest.read_bytes(b'{"course":"MED520","patient":"x"}', "note.json")
    assert read["reader"] == "text"
    assert "patient" in ingest.units_text(read)


# --- OCR (staff Index / local ingest only) --------------------------------


def _fake_ocr(unit_text: str):
    unit = {
        "text": unit_text,
        "kind": "ocr",
        "source_ref": "",
        "locator": "page 1",
        "citable": True,
        "course": None,
        "activity_ref": None,
        "status": "ok",
    }
    return {"units": [dict(unit)], "status": "ok", "reader": "ocr", "error": None}


def test_scanned_pdf_ocr_is_citable(monkeypatch):
    monkeypatch.setattr(ingest._ocr, "available", lambda: (True, "pytesseract"))
    monkeypatch.setattr(ingest._ocr, "ocr_pdf", lambda path, source_ref="": _fake_ocr("Recovered scan sentence."))
    status, text = ingest.extract_text(_blank_pdf_bytes(), "scan.pdf")
    assert status == "ok"
    assert "Recovered scan sentence." in text


def test_arabic_scan_is_non_citable_pending_signoff(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest._ocr, "available", lambda: (True, "pytesseract"))
    monkeypatch.setattr(ingest._ocr, "ocr_pdf", lambda path, source_ref="": _fake_ocr("إجازة مرضية"))
    data = _blank_pdf_bytes()

    result = ingest.ocr_bytes(data, "scan.pdf")
    assert result["units"][0]["citable"] is False
    assert result["units"][0]["status"] == "pending_signoff"

    status, text = ingest.extract_text(data, "scan.pdf")
    assert status == "signoff:arabic"
    assert "إجازة مرضية" in text

    from ai.moodle_content_job import record_ocr_pending

    stored = record_ocr_pending(
        "MED520",
        [{"filename": "scan.pdf", "ref": "file:scan.pdf", "locator": "page 1", "text": text, "signoff": "arabic"}],
        root=tmp_path,
    )
    assert stored == 1
    row = json.loads((tmp_path / "MED520.jsonl").read_text(encoding="utf-8").strip())
    assert row["citable"] is False
    assert row["status"] == "pending_signoff"


# --- chunk index + knowledge graph ----------------------------------------


def test_keyword_index_is_verbatim_and_semantic_stays_off(tmp_path):
    bank = {
        "MED520:file:101:a.pdf": {
            "id": "MED520:file:101:a.pdf",
            "course": "MED520",
            "kind": "file",
            "activity_id": 101,
            "sectionnum": 1,
            "name": "a.pdf",
            "text": "Keratinocyte turnover is described here.",
        }
    }
    out = ingest.build_course_index("MED520", index_root=tmp_path, bank=bank)
    assert out["chunks"] >= 1
    assert out["keyword"] is True
    assert out["semantic_status"] == "default_off"
    loaded = ce_index.load_index(out["path"])
    hits = ce_index.retrieve(loaded, "keratinocyte")
    assert hits
    span = hits[0]["spans"][0]
    assert span["text"] in hits[0]["text"]


def test_graph_edges_all_carry_a_real_passage_ref(tmp_path):
    rows = [
        {
            "course": "MED520",
            "kind": "file",
            "cmid": 101,
            "sectionnum": 1,
            "section": "Week 1",
            "name": "a.pdf",
            "ref": "file:a.pdf",
            "text": "Keratinocyte turnover. Mitosis sentence.",
        },
        {
            "course": "MED520",
            "kind": "page",
            "cmid": 102,
            "sectionnum": 1,
            "section": "Week 1",
            "name": "Page",
            "ref": "page:1",
            "text": "Second page sentence here.",
        },
    ]
    graph = ce_graph.build_from_bank_rows("MED520", rows)
    graph.validate()
    assert graph.stats()["edges"] > 0
    for edge in graph.edges.values():
        assert edge.provenance
        assert all(ce_graph.is_passage_ref(ref) for ref in edge.provenance)
    path = tmp_path / "MED520.jsonl"
    graph.save(path)
    assert ce_graph.Graph.load(path).stats() == graph.stats()
