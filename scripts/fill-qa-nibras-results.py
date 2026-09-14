"""Fill QA checklist (Nibras sheet) with validator results.

Idempotent: sets Status/Severity/Tester/Date/Findings/Action for NB-BP-001..013.
"""
import openpyxl

PATH = "docs/QA-Checklist-DataTrust-Nibras.xlsx"
TESTER = "QA Validator"
DATE = "2026-09-13"

# (status, severity, findings, action)
RESULTS = {
    "NB-BP-001": (
        "Pass", "Minor",
        "F1 grid height 0px (P3), F2 drawer overlap (P2), F3 enum capture (P3) — all fixed & E2E re-verified.",
        "—",
    ),
    "NB-BP-002": (
        "Pass", "—",
        "F4 positions/org rendering issue — fixed & re-verified.",
        "—",
    ),
    "NB-BP-003": (
        "Pass", "—",
        "F5 empty-state Add dialog, F6 compute 500, F7 backend down, F8 zero-salary seed — all fixed; full lifecycle Draft→Compute→Computed→Validate→Validated→Commit→Committed verified (run id=1).",
        "—",
    ),
    "NB-BP-004": (
        "Pass", "—",
        "Was FAIL (F10 feature-gap); resolved via people/self_views.py; payslip self-service view verified.",
        "—",
    ),
    "NB-BP-005": (
        "Pass", "—",
        "F14 leave-status label verified; remaining findings resolved.",
        "—",
    ),
    "NB-BP-006": (
        "Pass", "—",
        "Loan status sync fixed; end-to-end verified.",
        "—",
    ),
    "NB-BP-007": (
        "Pass", "—",
        "Manager inbox + FSM transitions verified.",
        "—",
    ),
    "NB-BP-008": (
        "Pass", "—",
        "Clean pass (leave policy + attendance).",
        "—",
    ),
    "NB-BP-009": (
        "Pass", "—",
        "Clean pass (certifications + rotation).",
        "—",
    ),
    "NB-BP-010": (
        "Pass", "—",
        "F20 org_unit required (was 500 IntegrityError) fixed + E2E verified (corr id=27); F21 archive action added + verified (corr 23 archived).",
        "—",
    ),
    "NB-BP-011": (
        "Pass", "Cosmetic",
        "F22 role-assignment audit log written + endpoint /accounts/role-audit-logs/; F23 refetchTables 403 for HR-only user fixed (no dataschema/tables call post-fix). F24 residual (P3): role-audit-logs API omits 'actor' field.",
        "Deferred",
    ),
    "NB-BP-012": (
        "Pass", "—",
        "Compensation masked: emp_1001 → basic_salary stripped + /compensation/ 403; ahmed → 420.000 + 200. Audit: view_compensation events (latest id=45). civil_id blank in real data (structural client-side masking only).",
        "—",
    ),
    "NB-BP-013": (
        "Pass", "—",
        "MDM 33 GOFSCO org units; DQ 16 people-domain rules; audit /catalog/governance-events/ = 45 events; no AASTMT/carbon leak (brand=nibras isolated).",
        "—",
    ),
}

wb = openpyxl.load_workbook(PATH)
ws = wb["Nibras"]

COL = {"status": 8, "severity": 10, "tester": 11, "date": 12, "findings": 13, "action": 14}

filled = 0
for r in range(3, ws.max_row + 1):
    bid = ws.cell(row=r, column=1).value
    if bid not in RESULTS:
        continue
    status, severity, findings, action = RESULTS[bid]
    ws.cell(row=r, column=COL["status"]).value = status
    ws.cell(row=r, column=COL["severity"]).value = severity
    ws.cell(row=r, column=COL["tester"]).value = TESTER
    ws.cell(row=r, column=COL["date"]).value = DATE
    ws.cell(row=r, column=COL["findings"]).value = findings
    ws.cell(row=r, column=COL["action"]).value = action
    filled += 1

wb.save(PATH)
print(f"Filled {filled} rows in '{PATH}' (Nibras sheet).")

# Verify readback
wb2 = openpyxl.load_workbook(PATH)
ws2 = wb2["Nibras"]
for r in range(3, 16):
    bid = ws2.cell(row=r, column=1).value
    st = ws2.cell(row=r, column=8).value
    sev = ws2.cell(row=r, column=10).value
    print(f"{bid}: status={st!r} severity={sev!r}")
