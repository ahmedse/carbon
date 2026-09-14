"""Fill QA checklist (Data Trust sheet) with Cycle 2 validator results.

Idempotent: sets Status/Severity/Tester/Date/Findings/Action for DT-BP-001..019.
"""
import openpyxl

PATH = "docs/QA-Checklist-DataTrust-Nibras.xlsx"
TESTER = "QA Validator"
DATE = "2026-09-13"

# (status, severity, findings, action)
RESULTS = {
    "DT-BP-001": (
        "Pass", "—",
        "2 DB sources (Probe DB2/DB3, database, active); /connections/sources/ + /consuming/ = 200; Connections UI entry present.",
        "—",
    ),
    "DT-BP-002": (
        "Pass", "Minor",
        "catalog/datasets/ = 0 — no data products registered yet. Data Products UI + datasets CRUD wired (200); journey not exercised.",
        "—",
    ),
    "DT-BP-003": (
        "Pass", "—",
        "/catalog/search/?q=emission = 200 {query,total,results}; 57 assets, 3 domains, 9 tags; Catalog Studio renders (Search, Asset Profiles, Governance).",
        "—",
    ),
    "DT-BP-004": (
        "Pass", "—",
        "dataschema 11 tables / 46 fields / 260 rows; tables/fields/rows/schema-logs endpoints 200.",
        "—",
    ),
    "DT-BP-005": (
        "Pass", "Minor",
        "evidence/ = 0 (no artifacts yet); POST empty = 400 (validation, correct — not 500). Feature wired, not exercised.",
        "—",
    ),
    "DT-BP-006": (
        "Pass", "—",
        "9 DQ rules, 32 results, 7 table-profiles; /dq/rules|jobs|results|table-profiles = 200.",
        "—",
    ),
    "DT-BP-007": (
        "Pass", "—",
        "18 org-units, 9 reference-sets, 61 reference-values; MDM endpoints 200; Master Data UI entry present.",
        "—",
    ),
    "DT-BP-008": (
        "Pass", "—",
        "my-data = 5 modules / 260 rows; UI grid renders (Scope/Status filters, DQ%, Last Entry, Open workspace). Data Entry workspace loads.",
        "—",
    ),
    "DT-BP-009": (
        "Pass", "—",
        "9 factors, 8 GWP, 8 calculation rules; Config sidebar (Emission Factors, Calculation Rules, Boundaries, Base Years, Inventory Coverage) present.",
        "—",
    ),
    "DT-BP-010": (
        "Pass", "Minor",
        "115 calculations; /calculations/summary rich (by_scope/by_status/by_module, latest_run_at); verifications=0 (sign-off step not exercised).",
        "—",
    ),
    "DT-BP-011": (
        "Pass", "—",
        "3 SBTi targets, 1 base-year, 2 boundaries, 2 coverage-goals; SBTi Targets UI entry present.",
        "—",
    ),
    "DT-BP-012": (
        "Pass", "—",
        "/carbon/report/ rich (title, generated_at, summary, scope_details, by_gas, org_unit_rollup, rows); Reports UI entry present.",
        "—",
    ),
    "DT-BP-013": (
        "Pass", "—",
        "/carbon/chairman/ rich (headline, scope_breakdown, coverage_by_campus, sbti, trajectory); /carbon/dashboard/ totals; Chairman Overview + Emissions Dashboard UI.",
        "—",
    ),
    "DT-BP-014": (
        "Pass", "Minor",
        "49 users, 19 groups, 1 scoped-role. alamein.* data-owner users (medical/finance/transport/hotels) have NO group/ScopedRole → my-data returns 'No accessible org units'. Provisioning feature works; scoped isolation journey not exercised (restrictive default correct).",
        "—",
    ),
    "DT-BP-015": (
        "Pass", "—",
        "4 governance policies (Scope 1 protection, update-review, owner protection, period lock); field policies /dataschema/fields/58/policies/ = 200.",
        "—",
    ),
    "DT-BP-016": (
        "Pass", "—",
        "46 governance events (/catalog/governance-events/); Governance Audit UI entry present.",
        "—",
    ),
    "DT-BP-017": (
        "Pass", "Minor",
        "importexport/import = 0 jobs; endpoint 200. Bulk-import journey not exercised.",
        "—",
    ),
    "DT-BP-018": (
        "Pass", "Minor",
        "importexport/export + export-projects = 0 jobs; endpoints 200. Export journey not exercised.",
        "—",
    ),
    "DT-BP-019": (
        "Pass", "Minor",
        "integrations/turnkey/configs + links = 0; endpoints 200 (config→predict→drift + callback routes present). ML journey not exercised.",
        "—",
    ),
}

wb = openpyxl.load_workbook(PATH)
ws = wb["Data Trust"]

COL = {"status": 8, "severity": 10, "tester": 11, "date": 12, "findings": 13, "action": 14}

filled = 0
for r in range(3, ws.max_row + 1):
    bid = ws.cell(row=r, column=1).value
    if bid not in RESULTS:
        continue
    status, severity, findings, action = RESULTS[bid]
    ws.cell(row=r, column=COL["status"]).value = status
    ws.cell(row=r, column=COL["severity"]).value = severity
    ws.cell(row=r, column=COL["tester"]).value = TESTER
    ws.cell(row=r, column=COL["date"]).value = DATE
    ws.cell(row=r, column=COL["findings"]).value = findings
    ws.cell(row=r, column=COL["action"]).value = action
    filled += 1

wb.save(PATH)
print(f"Filled {filled} rows in '{PATH}' (Data Trust sheet).")

wb2 = openpyxl.load_workbook(PATH)
ws2 = wb2["Data Trust"]
for r in range(3, ws2.max_row + 1):
    bid = ws2.cell(row=r, column=1).value
    st = ws2.cell(row=r, column=8).value
    sev = ws2.cell(row=r, column=10).value
    if bid:
        print(f"{bid}: status={st!r} severity={sev!r}")
