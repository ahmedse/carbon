#!/usr/bin/env python3
"""
Generate the QA module checklist workbook (Data Trust + Nibras) for the Carbon
Data Trust Platform.

Reproducible single source: edit the DATA_TRUST_ROWS / NIBRAS_ROWS lists below,
then run:

    .venv/bin/python scripts/generate-qa-checklist.py

Output: docs/QA-Checklist-DataTrust-Nibras.xlsx.

Columns (per checklist sheet):
  ID · Process Group · Business Process (End-to-End) · Role / Persona ·
  Journey Steps · Entry Route(s) · Backend Apps · Status · Priority ·
  Severity · Tester · Date Reviewed · Findings/Notes · Action

Pulse / AI console is intentionally EXCLUDED (still under development).
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT = REPO_ROOT / "docs" / "QA-Checklist-DataTrust-Nibras.xlsx"

# ── Columns ─────────────────────────────────────────────────────────────────
HEADERS = [
    "ID", "Process Group", "Business Process (End-to-End)", "Role / Persona",
    "Journey Steps (start → finish)", "Entry Route(s)", "Backend Apps",
    "Status", "Priority", "Severity", "Tester", "Date Reviewed",
    "Findings / Notes", "Action",
]

STATUS_OPTIONS = ["Not Started", "In Review", "Pass", "Fail", "Blocked", "N/A"]
PRIORITY_OPTIONS = ["P0 Critical", "P1 High", "P2 Medium", "P3 Low"]
SEVERITY_OPTIONS = ["—", "Blocker", "Major", "Minor", "Cosmetic"]
ACTION_OPTIONS = ["—", "Fix Required", "Accepted", "Deferred", "Re-Test"]

# ── Data: (Group, Business Process, Role, Journey Steps, Entry Routes, Backend Apps, Priority) ──
# Each row = ONE end-to-end business process performed by a relevant user role.
DATA_TRUST_ROWS = [
    # ── Data Lifecycle ──
    ("Data Lifecycle", "Onboard a new data source", "Data Connection Admin / Data Steward",
     "1. Open Connections and add a source (type, host, credentials).\n2. Test connectivity and confirm reachability.\n3. Discover schema/tables exposed by the source.\n4. Register selected tables to the Catalog.\n5. Tag metadata (domain, owner, sensitivity) and publish.",
     "/connections → /connections/:id → /catalog/metadata", "connections, dataschema, catalog", "P1 High"),
    ("Data Lifecycle", "Register & govern a data product", "Data Steward / Catalog Owner",
     "1. Create a Data Product in the Catalog.\n2. Attach schema/tables to the product.\n3. Assign owner, domain and tags.\n4. Define governance and access policy.\n5. Publish; confirm it is discoverable in search.",
     "/catalog/products → /catalog/products/:id → /catalog/governance", "catalog, dataschema", "P0 Critical"),
    ("Data Lifecycle", "Discover, request & get access to a dataset", "Data Consumer / Analyst",
     "1. Search the Catalog by keyword/domain/tag.\n2. Open asset detail and review metadata + schema.\n3. Request access.\n4. (Approver) grant access from Access Control.\n5. Confirm the consumer can now use the asset.",
     "/catalog/search → /catalog/assets/:id → /admin/access", "catalog, accounts", "P1 High"),
    ("Data Lifecycle", "Author / manage dataset schema & tables", "Data Schema Admin",
     "1. Open schema authoring for a data product.\n2. Add/edit tables and fields (types, constraints).\n3. Inspect row-level data detail.\n4. Validate the schema and publish.",
     "/catalog/products → /catalog/tables/:tableId → /carbon/my-data/row/:tableId/:rowId", "dataschema, catalog", "P1 High"),
    ("Data Lifecycle", "Attach & verify evidence (audit trail)", "Data Steward / Auditor",
     "1. Upload an evidence artifact.\n2. Link it to a record/asset.\n3. Verify the evidence and confirm it is traceable in the audit trail.",
     "/carbon-api/evidence/", "evidence, core", "P2 Medium"),

    # ── Data Quality & MDM ──
    ("Data Quality & MDM", "Run a data quality cycle", "Data Quality Analyst",
     "1. Open the DQ workspace and select a dataset.\n2. Define/attach a DQ rule.\n3. Run profiling and review scores.\n4. Triage failed issues (accept/fix).\n5. Close the loop and confirm the score improves.",
     "/dq → /dq/rules/:id", "dq, dataschema", "P0 Critical"),
    ("Data Quality & MDM", "Manage master / reference data (golden records)", "MDM Steward",
     "1. Open MDM reference data.\n2. Create/update a master record.\n3. Resolve duplicates (dedupe/merge).\n4. Confirm propagation to org units / consumers.",
     "/mdm → /admin/org-units → /admin/org-units/:id", "mdm", "P0 Critical"),

    # ── Carbon Footprint ──
    ("Carbon Footprint", "Enter activity data (Scope 1/2/3)", "Data Owner / Sustainability Officer",
     "1. Open My Data and pick a module.\n2. Select the table for the reporting period.\n3. Enter/import activity data.\n4. Validate and save the records.",
     "/carbon/my-data → /carbon/my-data/:moduleId → /carbon/my-data/:moduleId/:tableId", "emissions", "P0 Critical"),
    ("Carbon Footprint", "Configure emission factors, GWP & calculation rules", "Carbon Admin",
     "1. Manage emission factors (add/import).\n2. Maintain GWP reference values.\n3. Define calculation rules.\n4. Set boundaries, base years and inventory coverage.",
     "/carbon/admin/factors → /carbon/admin/gwp → /carbon/admin/rules → /carbon/admin/boundaries → /carbon/admin/base-years → /carbon/admin/inventory-coverage", "emissions", "P0 Critical"),
    ("Carbon Footprint", "Run & verify emissions calculations", "Carbon Admin / Verifier",
     "1. Trigger calculations for the reporting period.\n2. Review the calculated results.\n3. Verify the data and record sign-off.",
     "/carbon/calculations → /carbon/verification", "emissions", "P0 Critical"),
    ("Carbon Footprint", "Set SBTi targets & track progress", "Sustainability Officer / Carbon Admin",
     "1. Define SBTi targets against a base year.\n2. Confirm boundaries and coverage.\n3. Track progress against reporting periods.",
     "/carbon/admin/targets → /carbon/admin/base-years → /carbon/reporting/periods", "emissions", "P1 High"),
    ("Carbon Footprint", "Generate & publish emissions reports", "Analyst / Carbon Admin",
     "1. Open Reports and select a reporting period.\n2. Generate the report.\n3. Review figures, save and export.\n4. Confirm it surfaces for executive review.",
     "/carbon/reporting → /carbon/reporting/generate → /carbon/reporting/saved", "emissions, importexport", "P1 High"),
    ("Carbon Footprint", "Chairman / executive review (KPIs)", "Chairman / Executive",
     "1. Open Chairman overview for headline KPIs.\n2. Drill into the dashboard and analytics.\n3. Validate trends and drill-downs.",
     "/carbon/chairman → /carbon/dashboard → /carbon/analytics", "emissions", "P0 Critical"),

    # ── Access, Governance & Audit ──
    ("Access & Governance", "Provision a user & assign roles/capabilities", "Platform Administrator",
     "1. Create a user account.\n2. Assign to a group.\n3. Configure role matrix / capabilities.\n4. Verify access control (allowed vs denied).\n5. Confirm login and that only permitted areas are reachable.",
     "/admin/users → /admin/groups → /admin/role-matrix → /admin/access", "accounts", "P0 Critical"),
    ("Access & Governance", "Configure governance & field-level policies", "Governance Lead / Data Steward",
     "1. Define a governance policy.\n2. Set field-level policies (sensitivity, masking).\n3. Confirm policies are enforced in catalog/detail views.",
     "/admin/governance → /admin/catalog/field-policies", "catalog, dataschema", "P1 High"),
    ("Access & Governance", "Review the audit log", "Platform Admin / Auditor",
     "1. Open the audit log.\n2. Filter by actor/action/time.\n3. Investigate a specific change and confirm traceability.",
     "/admin/audit", "core", "P1 High"),

    # ── Data Movement & Integration ──
    ("Data Movement", "Import bulk data", "Import/Export Operator",
     "1. Open the Import/Export hub.\n2. Create an import job (file, mapping).\n3. Validate the import.\n4. Commit and confirm records land in the target.",
     "/import-export → /imports", "importexport", "P1 High"),
    ("Data Movement", "Export data / reports", "Import/Export Operator / Analyst",
     "1. Create an export job.\n2. Select scope/filters.\n3. Run and download the export.\n4. Verify the file contents.",
     "/import-export → /exports", "importexport", "P2 Medium"),
    ("Data Movement", "TurnKey ML integration (config → predict → drift)", "Integrations Admin / Data Scientist",
     "1. Create a TurnKey config.\n2. Create a link and register predictions.\n3. Review predictions and submit feedback.\n4. Check drift alerts and the webhook callback.",
     "/carbon-api/integrations/turnkey/*", "integrations", "P2 Medium"),
]

NIBRAS_ROWS = [
    # ── Employee Lifecycle ──
    ("Employee Lifecycle", "Onboard a new employee", "HR Admin",
     "1. Create the employee record.\n2. Assign position and org unit.\n3. Set compensation details.\n4. Capture civil ID and validate it.\n5. Activate the linked user account and confirm login.",
     "/people/employees → /people/employees/:id → /people/positions → /admin/org-units", "people, accounts, mdm", "P0 Critical"),
    ("Employee Lifecycle", "Manage positions & org structure", "HR Admin / MDM Steward",
     "1. Create/update positions.\n2. Map positions to org units.\n3. Handle rotation/reassignment.\n4. Confirm the org chart reflects changes.",
     "/people/positions → /people/rotation → /admin/org-units", "people, mdm", "P1 High"),

    # ── Payroll ──
    ("Payroll", "Run a payroll cycle", "Payroll Officer",
     "1. Open payroll runs and select a period.\n2. Run the calculation (salary + allowances − deductions).\n3. Review computed totals.\n4. Approve and generate payslips.",
     "/people/payroll → /people/payslip", "people", "P0 Critical"),
    ("Payroll", "Employee views payslip (self-service)", "Employee",
     "1. Log in to the employee portal.\n2. Open the My dashboard.\n3. View and download the payslip.",
     "/my → /people/payslip", "people", "P0 Critical"),

    # ── Leave & Requests ──
    ("Leave & Requests", "Employee requests leave (submit → approve)", "Employee / Manager",
     "1. Employee opens My Leave.\n2. Submits a leave request (type, dates).\n3. Manager reviews in the Team inbox and approves/rejects.\n4. Employee receives the outcome and the balance updates.",
     "/my/leave → /team/:id", "people, correspondence", "P1 High"),
    ("Leave & Requests", "Employee requests a loan (submit → approve → payroll)", "Employee / Manager / HR",
     "1. Employee submits a loan request from My Requests.\n2. Manager approves/rejects in the Team inbox.\n3. HR processes the loan.\n4. Payroll reflects the deduction in the next cycle.",
     "/my/requests → /team/:id → /people/loans → /people/payroll", "people, correspondence", "P2 Medium"),
    ("Leave & Requests", "Manager reviews & approves requests (inbox / FSM)", "Manager / Approver",
     "1. Open the Team approvals inbox.\n2. Open a request detail.\n3. Approve / reject / send back.\n4. Confirm the FSM advances and notifications fire.",
     "/team → /team/:id", "correspondence", "P0 Critical"),
    ("Leave & Requests", "Manage leave policy & attendance", "HR Admin",
     "1. Define/update the leave policy.\n2. Record attendance.\n3. Reconcile attendance with leave and payroll.",
     "/people/policies/:id → /people/attendance → /people/leave", "people", "P1 High"),

    # ── Other HR ──
    ("Other HR", "Track certifications & rotation", "HR Admin",
     "1. Record an employee certification.\n2. Schedule/manage rotation.\n3. Confirm status reflects in employee detail.",
     "/people/certifications → /people/rotation → /people/employees/:id", "people", "P2 Medium"),

    # ── Correspondence (e-office) ──
    ("Correspondence", "e-office correspondence workflow", "Correspondence Admin / Approver",
     "1. Create a letter/correspondence.\n2. Route to the approver.\n3. Approve and send.\n4. Archive and confirm it appears in the audit trail.",
     "/carbon-api/correspondence/ (+ Team inbox)", "correspondence, core", "P1 High"),

    # ── Access & Security ──
    ("Access & Security", "Provision HR roles & access (RBAC)", "Platform Admin / HR Admin",
     "1. Create a user and assign an HR group/role.\n2. Configure capabilities (people:manage vs people:view).\n3. Verify access and that the user lands in the right app.",
     "/admin/users → /admin/groups → /admin/access", "accounts, people", "P0 Critical"),
    ("Access & Security", "Verify sensitivity / data masking", "Auditor / Security",
     "1. Open employee detail as a non-authorized user.\n2. Confirm compensation is masked.\n3. Confirm civil ID is masked.\n4. Check the audit log for the access.",
     "/people/employees/:id → /admin/audit", "people, core", "P0 Critical"),
    ("Access & Security", "Verify shared platform (MDM / DQ / Audit) on Nibras", "Platform Admin",
     "1. Open MDM reference data.\n2. Open the DQ workspace.\n3. Review the audit log.\n4. Confirm brand isolation (Nibras data, not AASTMT).",
     "/mdm → /dq → /admin/audit", "mdm, dq, core, accounts", "P1 High"),
]

# ── Styling ─────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1e293b")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
AREA_FILL = PatternFill("solid", fgColor="dbeafe")
AREA_FONT = Font(bold=True, color="1e3a8a", size=11)
BAND_FILL = PatternFill("solid", fgColor="f8fafc")
THIN = Side(style="thin", color="e2e8f0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top")

COL_WIDTHS = [11, 20, 32, 24, 62, 34, 20, 13, 13, 12, 13, 14, 40, 13]


def build_sheet(ws, rows, title, id_prefix):
    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(bold=True, size=14, color="0f172a")
    t.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    # Headers
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[2].height = 30

    # Data
    r = 3
    idx = 0
    for group, process, role, steps, routes, apps, priority in rows:
        idx += 1
        row_id = f"{id_prefix}-BP-{idx:03d}"
        values = [row_id, group, process, role, steps, routes, apps, "Not Started", priority, "—", "", "", "", "—"]
        band = BAND_FILL if idx % 2 == 0 else None
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = WRAP
            if band:
                cell.fill = band
            if c in (8, 9, 10, 14):  # dropdown columns centered
                cell.alignment = CENTER
        lines = str(steps).count("\n") + 1
        ws.row_dimensions[r].height = max(30, lines * 14 + 6)
        r += 1

    last = r - 1

    # Freeze header + title
    ws.freeze_panes = "A3"

    # Auto filter
    ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{last}"

    # Column widths
    for c, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # Data validation dropdowns
    def dv(col, options):
        d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(d)
        d.add(f"{get_column_letter(col)}3:{get_column_letter(col)}{last}")
        return d

    dv(8, STATUS_OPTIONS)
    dv(9, PRIORITY_OPTIONS)
    dv(10, SEVERITY_OPTIONS)
    dv(14, ACTION_OPTIONS)


def build_readme(ws):
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110

    ws.merge_cells("A1:C1")
    ws["A1"] = "QA Business Process Checklist — Data Trust & Nibras"
    ws["A1"].font = Font(bold=True, size=16, color="0f172a")

    ws.merge_cells("A2:C2")
    ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · Pulse / AI console intentionally excluded (still under development)"
    ws["A2"].font = Font(italic=True, size=10, color="64748b")

    rows = [
        ("", "", ""),
        ("HOW TO USE", "", ""),
        ("", "1. Act as the relevant user", "Each row is ONE end-to-end business process. Log in as the Role / Persona listed, follow the Journey Steps start-to-finish, then record results. Do NOT review pages in isolation."),
        ("", "2. Status meanings", "Not Started → In Review → Pass | Fail | Blocked | N/A (not applicable to this instance)."),
        ("", "3. Severity (on Fail only)", "Blocker = cannot ship · Major = core path broken · Minor = edge case · Cosmetic = polish only."),
        ("", "4. Action", "Fix Required (assign to a worker) · Accepted (as-is) · Deferred (backlog) · Re-Test (after a fix)."),
        ("", "5. Priority", "P0 = auth/RBAC/payroll/data-integrity + primary business flows. P1 = main secondary flows. P2/P3 = long-tail & polish."),
        ("", "6. Journey Steps", "The numbered flow under 'Journey Steps' is the acceptance path — pass only if every step works in sequence for the listed persona."),
        ("", "", ""),
        ("SHEETS", "", ""),
        ("", "Data Trust", "AASTMT brand (aastmt) — the governed data + carbon platform. 19 end-to-end processes across Data Lifecycle, Data Quality & MDM, Carbon Footprint, Access/Governance/Audit, and Data Movement & Integration."),
        ("", "Nibras", "Nibras brand (nibras) — the HR & payroll platform. 13 end-to-end processes across Employee Lifecycle, Payroll, Leave & Requests, Correspondence (e-office), and Access & Security."),
        ("", "", ""),
        ("EXCLUDED (still under dev)", "", ""),
        ("", "Pulse / AI console", "All /admin/ai/* panels (Overview, Workspace, Conversations, Knowledge, Memory, Graph, Agents, Tools, Skills, Learning, Feedback, Monitoring, Audit). Review later when the AI console ships."),
    ]
    r = 4
    for a, b, c in rows:
        ca = ws.cell(row=r, column=1, value=a)
        cb = ws.cell(row=r, column=2, value=b)
        cc = ws.cell(row=r, column=3, value=c)
        if a in ("HOW TO USE", "SHEETS", "EXCLUDED (still under dev)"):
            ca.font = Font(bold=True, size=12, color="1d4ed8")
        if b and not b.startswith(" "):
            cb.font = Font(bold=True)
        cc.alignment = WRAP
        r += 1


def main():
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    build_readme(ws_readme)

    ws_dt = wb.create_sheet("Data Trust")
    build_sheet(ws_dt, DATA_TRUST_ROWS, "Data Trust — End-to-End Business Processes", "DT")

    ws_nb = wb.create_sheet("Nibras")
    build_sheet(ws_nb, NIBRAS_ROWS, "Nibras — End-to-End Business Processes", "NB")

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  Data Trust processes: {len(DATA_TRUST_ROWS)}")
    print(f"  Nibras processes:     {len(NIBRAS_ROWS)}")


if __name__ == "__main__":
    main()
