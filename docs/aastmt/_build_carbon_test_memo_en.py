#!/usr/bin/env python3
"""English memo for the AASTMT President. Ends: live dashboard and official report."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "2026-09-30-carbon-test-memo-en.docx"
ACCENT = RGBColor(0x00, 0x5C, 0xA8)
NAVY = RGBColor(0x1A, 0x1A, 0x2E)
MUTED = RGBColor(0x55, 0x55, 0x55)
LIGHT = "EEF4FB"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(table, color="CCCCCC"):
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{side}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:space"), "0")
                border.set(qn("w:color"), color)
                tcBorders.append(border)
            tcPr.append(tcBorders)


def run(paragraph, text, *, bold=False, size=11, color=NAVY, italic=False):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return r


def para(doc, text, *, bold=False, size=11, color=NAVY, before=0, after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run(p, text, bold=bold, size=size, color=color, italic=italic)
    return p


def rich(doc, parts, *, size=11, before=0, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for text, bold in parts:
        run(p, text, bold=bold, size=size)
    return p


def heading(doc, text):
    p = para(doc, text, bold=True, size=11, color=ACCENT, before=13, after=5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "005CA8")
    pBdr.append(bottom)
    pPr.append(pBdr)


def cell(cell_, text, *, size=10, bold=False, fill=None):
    p = cell_.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if fill:
        shade(cell_, fill)
    run(p, text, bold=bold, size=size, color=ACCENT if fill else NAVY)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    para(doc, "INTERNAL MEMORANDUM", bold=True, size=18, color=ACCENT, after=2)
    para(
        doc,
        "Arab Academy for Science, Technology & Maritime Transport",
        size=10,
        color=MUTED,
        after=10,
    )

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(
        (
            ("To", "Prof. Ismail Abdel Ghafar, President of the Academy"),
            ("From", "Dr. Ahmed Saied"),
            ("Date", "30 September 2026"),
            ("Subject", "The carbon platform is open — live dashboard and path to the official report"),
        )
    ):
        a, b = meta.rows[i].cells
        a.width = Cm(2.5)
        b.width = Cm(13.5)
        cell(a, label, bold=True, fill=LIGHT)
        cell(b, value, bold=(i == 3))
    borders(meta)

    para(doc, "Sir,", before=16, after=10)

    para(
        doc,
        "The carbon platform the Academy has been waiting for is now in the light.",
    )

    heading(doc, "What now stands")
    rich(
        doc,
        [
            ("The ", False),
            ("Data Trust Platform", True),
            (
                " is the governed place for Academy data. A campus sees its own sources. The centre sees the joined picture. Nothing reaches a report unless it has a source, a check, and a named owner.",
                False,
            ),
        ],
    )
    rich(
        doc,
        [
            ("Carbon Footprint", True),
            (
                " is the first application on that place. Campuses enter electricity, fuel, and travel. Two ends sit on the same record.",
                False,
            ),
        ],
    )
    rich(
        doc,
        [
            ("The live dashboard is ", False),
            ("Chairman Overview", True),
            (
                ". It shows what the platform has already joined. That is the light for this period.",
                False,
            ),
        ],
    )
    rich(
        doc,
        [
            (
                "The official report is the Academy’s published footprint. It will be issued from this same record, when the President so decides — not from a file assembled for a single occasion.",
                False,
            ),
        ],
    )
    rich(
        doc,
        [
            ("Pulse", True),
            (
                " is the coworker inside the platform. A question about a line on the dashboard is answered from the same data on the screen. Pulse does not invent data. Pulse does not write. A change to the record leaves conversation: a plan, an approval, a confirmed step.",
                False,
            ),
        ],
    )

    heading(doc, "What this period is")
    para(
        doc,
        "The dashboard is open. The official report is not yet issued. That stamp remains with the President, and is given when the record is ready. Until then the live view should not be quoted in a speech, a letter, or a paper for the Board. The page itself will not pretend otherwise.",
    )

    heading(doc, "The request")
    para(
        doc,
        "Open Carbon Footprint, then Chairman Overview, and say whether that live dashboard is fit for a president. If a line is obscure, ask Pulse.",
    )
    para(
        doc,
        "Eng. Mostafa Kamel and Dr. Mostafa Saad will take the campuses from here. The President’s account is for reading. Sign-in will be handed over. It is not written here.",
    )

    para(doc, "Respectfully,", before=16, after=4)
    para(doc, "Dr. Ahmed Saied", bold=True, after=12)
    para(
        doc,
        "Internal. Do not attach the sign-in to this memorandum.",
        size=9,
        color=MUTED,
        italic=True,
        after=0,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
