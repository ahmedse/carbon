#!/usr/bin/env python3
"""Generate Nibras customer handoff packs: Word guides + Excel UAT workbook."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent
VERSION = "v0.1"
RELEASE_DATE = date.today().isoformat()
PRODUCT = "Nibras (People · My · Team)"


def _set_narrow_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)


def _heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def _para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def _bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def _cover(doc, title, subtitle, audience):
    _set_narrow_margins(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(PRODUCT)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0xB8, 0x5A, 0x1A)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(22)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitle)
    r.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Audience: {audience}\nDocument version: {VERSION}\nRelease date: {RELEASE_DATE}\n"
        "Classification: Customer handoff — Internal use with customer"
    ).font.size = Pt(10)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        "This guide describes production behaviour for the deployed Nibras instance. "
        "Screens may vary slightly by role and configuration. For support, contact your "
        "Nibras administrator."
    ).italic = True


def build_my_guide():
    doc = Document()
    _cover(
        doc,
        "My — Employee Self-Service User Guide",
        "How to check leave, submit requests, and track status",
        "Employees (all staff with My access)",
    )

    _heading(doc, "1. What is My?", 1)
    _para(
        doc,
        "My is your personal HR self-service area. Use it to see your profile, leave balances, "
        "submit leave and profile-change requests, and track what you have submitted. "
        "Managers approve requests in Team; HR administers records in People.",
    )

    _heading(doc, "2. Before you start", 1)
    _bullets(
        doc,
        [
            "You need a Nibras login linked to your employee record (usually emp_<employee number>).",
            "Use the URL provided by your organisation for this deployment.",
            "Preferred browsers: current Chrome, Edge, or Safari.",
            "If My shows empty leave balances, ask HR — entitlements may not be assigned yet.",
        ],
    )

    _heading(doc, "3. Sign in", 1)
    _numbered(
        doc,
        [
            "Open the Nibras URL in your browser.",
            "Enter your username and password.",
            "After sign-in, open My from the main navigation.",
        ],
    )

    _heading(doc, "4. Dashboard (My home)", 1)
    _para(doc, "Path: My → Dashboard ( /my )")
    _bullets(
        doc,
        [
            "See your employee profile summary.",
            "Open leave, requests, and (where enabled) loans from this page.",
            "Use Request Profile Change when personal data needs an HR-approved update.",
        ],
    )

    _heading(doc, "5. View leave balance", 1)
    _para(doc, "Path: My → My Leave ( /my/leave )")
    _numbered(
        doc,
        [
            "Open My Leave.",
            "Review entitled, carried, used, pending, and remaining days by leave type.",
        ],
    )
    _para(
        doc,
        "Tip: Pending days are requests still waiting for manager approval. Remaining is what you can still request.",
        bold=False,
    )

    _heading(doc, "6. Request leave (day off)", 1)
    _numbered(
        doc,
        [
            "Go to My → My Leave.",
            "Click Request Leave.",
            "Choose leave type (e.g. annual, sick).",
            "Select start and end dates.",
            "Add a short note if required by your policy.",
            "Submit.",
        ],
    )
    _para(doc, "Expected result:", bold=True)
    _bullets(
        doc,
        [
            "You see a confirmation.",
            "The request appears under My Requests with status submitted / in review.",
            "Your manager receives it in Team → Approvals Inbox.",
        ],
    )

    _heading(doc, "7. Track my requests", 1)
    _para(doc, "Path: My → My Requests ( /my/requests )")
    _numbered(
        doc,
        [
            "Open My Requests to see all items you submitted.",
            "Click a request to see status and history.",
            "If a manager sends it back, edit and resubmit as instructed.",
        ],
    )

    _heading(doc, "8. Request a profile change", 1)
    _numbered(
        doc,
        [
            "From My Dashboard, choose Request Profile Change.",
            "Enter the fields you need updated.",
            "Submit for approval.",
        ],
    )
    _para(
        doc,
        "Profile changes are governed — they do not update your record until approved.",
    )

    _heading(doc, "9. My loans (if shown)", 1)
    _para(
        doc,
        "Where enabled, the dashboard lists loans issued to you. Contact HR for new loan requests "
        "or repayment questions unless your instance exposes a request form.",
    )

    _heading(doc, "10. Common problems", 1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Symptom", "What to do"),
        ("Cannot see My in the menu", "Ask HR/IT to link your user to an employee and grant My access."),
        ("Leave balance is zero", "Ask HR to propagate leave policies / assign entitlements for your year."),
        ("Request stuck as submitted", "Confirm your manager is set on your employee record; remind them to open Team."),
        ("Forgot password", "Use the organisation’s password reset process or contact your administrator."),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b

    _heading(doc, "11. Quick FAQ", 1)
    _bullets(
        doc,
        [
            "Who approves my leave? Your assigned manager in Team (and further steps if configured).",
            "Can I cancel leave? Follow the status on My Requests; if not available, ask your manager or HR.",
            "Where is payroll / payslip? Payslips are managed under People/Payroll by HR — ask HR for your payslip process.",
        ],
    )

    _heading(doc, "12. Document control", 1)
    _para(doc, f"Product: {PRODUCT}  |  Version: {VERSION}  |  Date: {RELEASE_DATE}")
    _para(doc, "Owner: Nibras delivery / customer success  |  Review: after each production release")

    path = OUT / f"Nibras-My-Employee-User-Guide-{VERSION}.docx"
    doc.save(path)
    return path


def build_people_guide():
    doc = Document()
    _cover(
        doc,
        "People — HR Administration Guide",
        "Day-to-day workforce, leave policy, payroll, and compliance operations",
        "HR officers / People administrators",
    )

    _heading(doc, "1. What is People?", 1)
    _para(
        doc,
        "People is the HR back office for Nibras: employees, organisation, leave policies, "
        "compensation, loans, payroll runs, and compliance. Employees use My; managers use Team.",
    )

    _heading(doc, "2. Recommended daily / weekly flow", 1)
    _numbered(
        doc,
        [
            "Set up / maintain Organisation and Positions.",
            "Onboard employees and assign managers (required for Team approvals).",
            "Maintain leave policies; propagate entitlements for the year.",
            "Monitor leave records and exceptions.",
            "Maintain compensation / benefits; verify ledger lines before payroll.",
            "Run payroll: create → compute → validate → commit → WPS export (when authorised).",
            "Keep compliance rules authoritative and versioned.",
        ],
    )

    _heading(doc, "3. Add an employee", 1)
    _para(doc, "Path: People → Employees")
    _numbered(
        doc,
        [
            "Click Add Employee.",
            "Enter name, civil ID / national ID, employee number, join/hire date.",
            "Select position, basic salary, contract type.",
            "Set Kuwaitization and rotation where applicable.",
            "Set Manager to the approving supervisor (critical for leave approvals).",
            "Save and confirm the employee appears in the list.",
        ],
    )

    _heading(doc, "4. Edit, deactivate, reactivate", 1)
    _numbered(
        doc,
        [
            "Open the employee detail; edit fields such as salary or position; save.",
            "Deactivate when the person leaves active duty; confirm they leave active lists.",
            "Reactivate when returning; confirm is_active is restored.",
        ],
    )

    _heading(doc, "5. Positions", 1)
    _para(doc, "Path: People → Positions — create/edit title, grade, and org unit so every employee is anchored.")

    _heading(doc, "6. Leave policies", 1)
    _para(doc, "Path: People → Policies")
    _bullets(
        doc,
        [
            "Create policies with leave type, entitled days, accrual, carryover, status, and effective dates.",
            "Use category and tags for grouping/filtering.",
            "Scope by Kuwaitization (e.g. Kuwaiti 42 vs expat 30) and rotation patterns where required.",
            "Use New Version / Fork for immutable history; review Version History.",
            "Propagate to entitlements (dry-run first, then apply) for the target year.",
        ],
    )

    _heading(doc, "7. Leave entitlements & records", 1)
    _para(
        doc,
        "Path: People → Leave — review entitlements and create manual leave records when HR must post leave outside My.",
    )

    _heading(doc, "8. Benefits & compensation", 1)
    _bullets(
        doc,
        [
            "Config → Benefit Types: ACCOM, VEHICLE, MEDICAL, SCHOOL, TICKETS, OT_BASE (as used by your org).",
            "Assign benefits to employees with values.",
            "Define compensation components and plans; verify per-employee ledger lines before compute.",
        ],
    )
    _para(
        doc,
        "Important: Do not run production payroll on estimated salaries until ledger data is verified (payroll freeze until salary data is production-safe).",
        bold=True,
    )

    _heading(doc, "9. Loans", 1)
    _para(doc, "Path: People → Loans — issue a loan and confirm installment schedule is generated.")

    _heading(doc, "10. Payroll run", 1)
    _para(doc, "Path: People → Payroll")
    _numbered(
        doc,
        [
            "Create a PayrollRun for the period.",
            "Compute payslip lines.",
            "Validate and resolve issues.",
            "Commit the run.",
            "Export WPS when required for banking.",
        ],
    )

    _heading(doc, "11. Compliance & EOSI", 1)
    _bullets(
        doc,
        [
            "Policies / Config → Compliance Rules: keep authoritative rules versioned with provenance.",
            "Employee → EOSI and Timeline for end-of-service indemnity and lifecycle events.",
        ],
    )

    _heading(doc, "12. Manager hierarchy (Team depends on this)", 1)
    _para(
        doc,
        "Import tools do not always set managers. Until Employee.manager is set, Team inboxes stay empty "
        "and leave will not route. After linking users, set managers in People → Employees (or via approved CSV process).",
    )

    _heading(doc, "13. Enabling My for staff (ops summary)", 1)
    _bullets(
        doc,
        [
            "Ensure correspondence / leave workflows are seeded for the instance.",
            "Import or create employees; propagate leave policies.",
            "Link users (emp_<employee_no>) with the instance password procedure.",
            "Assign managers; re-link if Team access must refresh.",
        ],
    )
    _para(
        doc,
        "Detailed command sequence for operators: see GOFSCO-ONBOARDING-RUNBOOK (internal ops). "
        "This user-facing guide does not replace that runbook.",
    )

    _heading(doc, "14. FAQ", 1)
    _bullets(
        doc,
        [
            "How do I add an employee? People → Employees → Add; assign position and required fields.",
            "How do I run payroll? People → Payroll → create run → compute → validate → commit → WPS.",
            "How is data protected? Role-based access: People Admin, org-scoped owners, read-only analysts as configured.",
        ],
    )

    _heading(doc, "15. Document control", 1)
    _para(doc, f"Product: {PRODUCT}  |  Version: {VERSION}  |  Date: {RELEASE_DATE}")

    path = OUT / f"Nibras-People-HR-Admin-Guide-{VERSION}.docx"
    doc.save(path)
    return path


def build_team_guide():
    doc = Document()
    _cover(
        doc,
        "Team — Manager Approvals Guide",
        "How to approve, reject, and send back team requests",
        "Managers / supervisors with Team access",
    )

    _heading(doc, "1. What is Team?", 1)
    _para(
        doc,
        "Team is the approvals inbox for people who manage others. Leave, loans, and profile-change "
        "requests from My appear here when you are the assigned manager (or next approver in the chain).",
    )

    _heading(doc, "2. Open the inbox", 1)
    _para(doc, "Path: Team → Approvals Inbox ( /team )")
    _bullets(
        doc,
        [
            "You only see Team if your account has manager access.",
            "Items typically show as submitted or in review.",
        ],
    )

    _heading(doc, "3. Approve a request", 1)
    _numbered(
        doc,
        [
            "Open a request from the inbox.",
            "Review dates, type, and employee note.",
            "Click Approve.",
        ],
    )
    _para(doc, "Expected: Status advances; employee sees approved status and balance updates in My.")

    _heading(doc, "4. Reject a request", 1)
    _numbered(
        doc,
        [
            "Open the request → Reject.",
            "Enter a clear reason.",
            "Confirm.",
        ],
    )

    _heading(doc, "5. Send back for revision", 1)
    _numbered(
        doc,
        [
            "Open the request → Send Back.",
            "Add a note describing what to fix.",
            "Employee edits and resubmits from My.",
        ],
    )

    _heading(doc, "6. Multi-step chains", 1)
    _para(
        doc,
        "Some requests need supervisor then HR. You only see the item when it is your turn. "
        "Final approval happens only after the last step.",
    )

    _heading(doc, "7. If the inbox is empty", 1)
    _bullets(
        doc,
        [
            "Confirm employees have you set as Manager in People.",
            "Confirm staff have submitted requests from My.",
            "Ask HR/IT to refresh Team access after manager changes.",
        ],
    )

    _heading(doc, "8. Document control", 1)
    _para(doc, f"Product: {PRODUCT}  |  Version: {VERSION}  |  Date: {RELEASE_DATE}")

    path = OUT / f"Nibras-Team-Manager-Guide-{VERSION}.docx"
    doc.save(path)
    return path


# ── Excel UAT ─────────────────────────────────────────────────────────────

THIN = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
HEADER_FILL = PatternFill("solid", fgColor="B85A1A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="F5E6D3")
PASS_FILL = PatternFill("solid", fgColor="E8F5E9")
FAIL_FILL = PatternFill("solid", fgColor="FFEBEE")


def _style_header(ws, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def _autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _add_result_validation(ws, result_col, start_row, end_row):
    dv = DataValidation(type="list", formula1='"PASS,FAIL,BLOCKED,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(result_col)}{start_row}:{get_column_letter(result_col)}{end_row}")


CASES = {
    "A_People": [
        ("A1", "Add employee", "HR creates employee with position, salary, Kuwaitization, rotation, manager",
         "People → Employees → Add; fill required fields; Save",
         "Row appears; detail GET persists all fields", "At least one Position exists"),
        ("A2", "Edit / deactivate / reactivate", "Lifecycle controls work",
         "Edit salary/position; Deactivate; Reactivate",
         "is_active and fields match after each step", "Employee from A1"),
        ("A3", "Positions CRUD", "Create and edit Position",
         "People → Positions → Add → Edit → Save",
         "Edits persist", ""),
        ("A4", "Compliance rules", "Authoritative rule with schema",
         "Policies/Config → Compliance Rules → Add authoritative rule",
         "Rule listed with is_authoritative and provenance", ""),
        ("A5", "Leave policy create", "Policy with category/tags",
         "People → Policies → Add Policy; set category/tags; Save",
         "Filter by category/tag works", ""),
        ("A6", "Leave policy scope", "Kuwaiti 42 / expat 30 / rotation scope",
         "Create scoped policies; Save",
         "Scope fields persist on detail", ""),
        ("A7", "Policy versioning", "Fork / version history",
         "New Version/Fork; open Version History",
         "≥2 versions; snapshot includes scope fields", "Active policy exists"),
        ("A8", "Propagate entitlements", "Dry-run then apply for year",
         "Policy detail → Propagate",
         "Entitlements only for eligible; idempotent re-run", "Scoped policies + employees"),
        ("A9", "Leave entitlements/records", "HR lists and posts leave",
         "People → Leave; add LeaveRecord",
         "Record listed", ""),
        ("A10", "Benefits", "Benefit types + assign",
         "Config → Benefit Types; assign to employee",
         "Types and benefits listed", ""),
        ("A11", "Compensation ledger verify", "Components, plan, verify line",
         "Config components/plan; verify employee ledger line",
         "Verify returns success", ""),
        ("A12", "Loans", "Issue loan + installments",
         "People → Loans → add",
         "Loan + installments listed", ""),
        ("A13", "Attendance", "Record + permission",
         "People → Attendance",
         "Rows persist", "Skip if hidden from go-live nav"),
        ("A14", "Certifications", "Add certification",
         "People → Certifications",
         "Shows on employee", ""),
        ("A15", "Rotation schedules", "Define pattern",
         "People → Rotation",
         "Schedule listed", "Skip if hidden from go-live nav"),
        ("A16", "Payroll lifecycle", "Compute → validate → commit → WPS",
         "People → Payroll full cycle",
         "Each step succeeds; WPS downloadable", "Verified ledger; salary production-safe"),
        ("A17", "EOSI + timeline", "Compute EOSI; view timeline",
         "Employee → EOSI; Timeline",
         "Numeric EOSI; events listed", ""),
    ],
    "B_My": [
        ("B1", "My profile", "Employee sees own profile",
         "My → Dashboard", "Own profile rendered", "User linked to employee"),
        ("B2", "Leave balance", "Entitled/carried/used/pending/remaining",
         "My → My Leave", "Balance widget shows five figures", "Entitlements propagated"),
        ("B3", "Request leave", "Submit leave request",
         "My Leave → Request Leave → submit",
         "Confirmation; appears in My Requests", "Manager assigned"),
        ("B4", "Track requests", "Status + history",
         "My → My Requests → open detail",
         "Status visible", "Request from B3"),
        ("B5", "Profile change request", "Submit personal-data change",
         "Dashboard → Request Profile Change",
         "Request created and routable", ""),
        ("B6", "My loans", "See own loans",
         "Dashboard loans section",
         "Own loans listed", "Loan exists for employee"),
    ],
    "C_Team": [
        ("C1", "Approvals inbox", "Manager sees actionable items",
         "Team → Approvals Inbox",
         "Submitted/in_review items appear", "Manager hierarchy set"),
        ("C2", "Approve", "Approve leave",
         "Open request → Approve",
         "Status approved; employee balance updates", "Item from B3"),
        ("C3", "Reject", "Reject with reason",
         "Open → Reject → reason",
         "Rejected with reason stored", ""),
        ("C4", "Send back", "Return for revision",
         "Open → Send Back → note",
         "Requestor can edit/resubmit", ""),
        ("C5", "Multi-step chain", "Two-step approval",
         "Approve as step-1 then step-2",
         "Final only after last step", "Workflow configured"),
    ],
    "D_Journey": [
        ("D1", "End-to-end leave loop", "My → Team → My → People",
         "1) Employee requests leave 2) Manager approves 3) Employee sees approved + balance 4) HR sees LeaveRecord",
         "No silent data drop across apps", "Two users: employee + manager"),
    ],
}


def _write_cases_sheet(wb, title, cases):
    ws = wb.create_sheet(title)
    headers = [
        "ID", "Case", "Goal", "Steps", "Expected / PASS criteria",
        "Preconditions", "Result", "Tester", "Date", "Evidence / notes",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))
    for row in cases:
        ws.append([*row, "", "", "", ""])
    for r in range(2, len(cases) + 2):
        for c in range(1, 11):
            cell = ws.cell(r, c)
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _add_result_validation(ws, 7, 2, len(cases) + 1)
    _autosize(ws, [6, 22, 36, 44, 40, 28, 12, 14, 12, 28])
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    return ws


def build_uat_workbook():
    wb = Workbook()

    # Cover
    cover = wb.active
    cover.title = "00_Cover"
    cover["A1"] = PRODUCT
    cover["A1"].font = Font(bold=True, size=16, color="B85A1A")
    cover["A2"] = "UAT / QA Checklist & Sign-off Workbook"
    cover["A2"].font = Font(bold=True, size=14)
    cover["A4"] = f"Document version: {VERSION}"
    cover["A5"] = f"Release date: {RELEASE_DATE}"
    cover["A6"] = "How to use"
    cover["A6"].font = Font(bold=True)
    cover["A7"] = (
        "1) Fill 01_Environment with the customer URLs and test accounts.\n"
        "2) Execute sheets A → B → C → D top to bottom.\n"
        "3) Set Result to PASS / FAIL / BLOCKED / N/A from the dropdown.\n"
        "4) Attach evidence (screenshot ID or ticket) in Evidence / notes.\n"
        "5) Complete 99_SignOff and obtain customer signature or email ACK."
    )
    cover["A7"].alignment = Alignment(wrap_text=True)
    cover.merge_cells("A7:F12")
    cover["A14"] = "Source of cases: docs/QA-MANUAL-PEOPLE-MY-TEAM.md (customer-facing rewrite)"
    cover["A15"] = "Related ops (internal): docs/nibras/GOFSCO-ONBOARDING-RUNBOOK.md"
    _autosize(cover, [100])

    # Environment
    env = wb.create_sheet("01_Environment")
    env.append(["Item", "Value (fill in)", "Notes"])
    _style_header(env, 3)
    for row in [
        ("Frontend URL", "", "Customer production or UAT URL"),
        ("Backend / API base", "", "If distinct from FE proxy"),
        ("Brand / instance", "nibras", ""),
        ("HR admin test user", "", "People access"),
        ("Employee test user", "", "My access — emp_<no>"),
        ("Manager test user", "", "Team access — manages employee above"),
        ("Browser", "", "Chrome / Edge / Safari + version"),
        ("UAT start date", "", ""),
        ("UAT end date", "", ""),
        ("Build / release tag", "", "Match deployed version"),
    ]:
        env.append(list(row))
    for r in range(1, 12):
        for c in range(1, 4):
            env.cell(r, c).border = THIN
    _autosize(env, [28, 40, 40])

    _write_cases_sheet(wb, "A_People", CASES["A_People"])
    _write_cases_sheet(wb, "B_My", CASES["B_My"])
    _write_cases_sheet(wb, "C_Team", CASES["C_Team"])
    _write_cases_sheet(wb, "D_CrossApp", CASES["D_Journey"])

    # GOFSCO matrix
    mx = wb.create_sheet("E_GOFSCO_Matrix")
    mx.append(["#", "Requirement", "Where", "Verify via", "Result", "Notes"])
    _style_header(mx, 6)
    for row in [
        (1, "Compliance rule engine (authoritative, versioned)", "People ComplianceRule", "A4", "", ""),
        (2, "Kuwaitization leave (42 vs 30)", "LeavePolicy scope", "A6 + A8", "", ""),
        (3, "Rotation leave scoping", "LeavePolicy rotations", "A6 + A8", "", ""),
        (4, "Payroll lifecycle → WPS", "PayrollRun", "A16", "", ""),
        (5, "EOSI indemnity", "Employee EOSI", "A17", "", ""),
        (6, "Benefits & compensation categories", "BenefitType / Comp", "A10 + A11", "", ""),
        (7, "HR Reporting module", "Not built", "N/A roadmap", "N/A", "Acknowledged gap"),
        (8, "Full HR module registry", "Not built", "N/A roadmap", "N/A", "Acknowledged gap"),
    ]:
        mx.append(list(row))
    _add_result_validation(mx, 5, 2, 9)
    for r in range(1, 10):
        for c in range(1, 7):
            mx.cell(r, c).border = THIN
            mx.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(mx, [4, 42, 24, 16, 12, 24])

    # Sign-off
    so = wb.create_sheet("99_SignOff")
    so["A1"] = "UAT Sign-off"
    so["A1"].font = Font(bold=True, size=14, color="B85A1A")
    so["A3"] = "Area"
    so["B3"] = "PASS?"
    so["C3"] = "Comments"
    for col in range(1, 4):
        so.cell(3, col).fill = HEADER_FILL
        so.cell(3, col).font = HEADER_FONT
        so.cell(3, col).border = THIN
    for i, area in enumerate(
        ["Part A — People", "Part B — My", "Part C — Team", "Part D — Cross-app journey",
         "GOFSCO matrix 1–6", "Known gaps 7–8 acknowledged"],
        start=4,
    ):
        so.cell(i, 1).value = area
        so.cell(i, 2).value = ""
        so.cell(i, 3).value = ""
        for c in range(1, 4):
            so.cell(i, c).border = THIN
    _add_result_validation(so, 2, 4, 9)

    so["A11"] = "Overall recommendation"
    so["A11"].font = Font(bold=True)
    so["B11"] = "READY / READY WITH EXCEPTIONS / NOT READY"
    so["A13"] = "Vendor tester name"
    so["B13"] = ""
    so["A14"] = "Vendor tester date"
    so["B14"] = ""
    so["A15"] = "Customer UAT lead name"
    so["B15"] = ""
    so["A16"] = "Customer UAT lead date"
    so["B16"] = ""
    so["A17"] = "Customer signature / email ACK reference"
    so["B17"] = ""
    so["A19"] = (
        "By signing, the customer confirms the executed cases reflect acceptance of People / My / Team "
        "for the stated release, subject to any FAIL/BLOCKED items listed as agreed exceptions."
    )
    so["A19"].alignment = Alignment(wrap_text=True)
    so.merge_cells("A19:C22")
    _autosize(so, [42, 36, 40])

    path = OUT / f"Nibras-UAT-QA-Checklist-{VERSION}.xlsx"
    wb.save(path)
    return path


def main():
    paths = [
        build_my_guide(),
        build_people_guide(),
        build_team_guide(),
        build_uat_workbook(),
    ]
    readme = OUT / "README.md"
    readme.write_text(
        f"""# Nibras customer handoff pack ({VERSION})

Generated: {RELEASE_DATE}

## Files

| File | Format | Audience |
|------|--------|----------|
| `Nibras-My-Employee-User-Guide-{VERSION}.docx` | Word | Employees |
| `Nibras-People-HR-Admin-Guide-{VERSION}.docx` | Word | HR / People admins |
| `Nibras-Team-Manager-Guide-{VERSION}.docx` | Word | Managers |
| `Nibras-UAT-QA-Checklist-{VERSION}.xlsx` | Excel | Customer QA + delivery lead |

## How to hand over

1. Confirm the **release tag / build** matches production; put it on the Excel Environment sheet.
2. Send the **three Word guides** for training (and optional PDF export from Word).
3. Run **UAT** with the Excel workbook on the customer tenant; collect Sign-off sheet.
4. Keep internal ops runbook (`../GOFSCO-ONBOARDING-RUNBOOK.md`) — do not send raw unless the customer ops team needs it.

## Regenerate

```bash
/home/ahmed/ws/carbon/.venv/bin/python /home/ahmed/ws/carbon/docs/nibras/handoff/generate_handoff_pack.py
```
""",
        encoding="utf-8",
    )
    for p in paths:
        print(p)


if __name__ == "__main__":
    main()
